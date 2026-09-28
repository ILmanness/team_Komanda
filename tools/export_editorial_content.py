"""Convert the supplied editorial workbooks into versioned application data.

Usage: python tools/export_editorial_content.py KNOWLEDGE.xlsx TRAINING.xlsx
This reads workbooks only; the JSON output can be reviewed and imported without Excel.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'backend' / 'content'


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    return value


def rows(sheet):
    return [[clean(cell) for cell in row] for row in sheet.values]


def table(sheet):
    values = rows(sheet)
    headers = values[0]
    return [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]


def sections_from_rows(sheet):
    sections = []
    current = None
    for row in range(5, sheet.max_row + 1):
        cell = sheet.cell(row, 2)
        value = clean(cell.value)
        if not value or value == 'Ваш ответ:':
            continue
        if cell.font.bold:
            current = {'title': str(value), 'paragraphs': []}
            sections.append(current)
        else:
            if current is None:
                current = {'title': 'Введение', 'paragraphs': []}
                sections.append(current)
            current['paragraphs'].append(str(value))
    return [section for section in sections if section['paragraphs']]


def quiz_from_sheet(sheet):
    introduction = str(sheet.cell(6, 2).value or '')
    short = re.search(r'короткой самопроверки ответьте на вопросы ([^.]+)', introduction)
    if not short:
        raise ValueError(f'{sheet.title}: short quiz selection is missing')
    short_numbers = [int(value) for value in re.findall(r'\d+', short.group(1))]
    questions = {}
    current = None
    answer_section = False
    for row in range(5, sheet.max_row + 1):
        value = clean(sheet.cell(row, 2).value)
        if not value:
            continue
        if value == 'Ответы и объяснения':
            answer_section = True
            current = None
            continue
        if not answer_section:
            match = re.fullmatch(r'Вопрос (\d+)\. Материал ([\d.]+)', str(value))
            if match:
                number = int(match.group(1))
                current = {'number': number, 'material_id': match.group(2),
                           'prompt': '', 'choices': [], 'correct': None, 'explanation': ''}
                questions[number] = current
            elif current and not current['prompt']:
                current['prompt'] = str(value)
            elif current and re.match(r'^[ABCD]\. ', str(value)):
                current['choices'].append({'id': value[0], 'text': value[3:]})
        else:
            match = re.fullmatch(r'Вопрос (\d+)\. Ответ ([ABCD])\. Вернуться к материалу ([\d.]+)', str(value))
            if match:
                number = int(match.group(1))
                if number not in questions or questions[number]['material_id'] != match.group(3):
                    raise ValueError(f'{sheet.title}: answer key mismatch')
                current = questions[number]
                current['correct'] = match.group(2)
            elif current and not current['explanation']:
                current['explanation'] = str(value)
                current = None
    if len(short_numbers) != 3 or len(questions) not in (5, 6, 9):
        raise ValueError(f'{sheet.title}: unexpected question count')
    if any(len(question['choices']) != 4 or not question['correct'] or not question['explanation']
           for question in questions.values()):
        raise ValueError(f'{sheet.title}: incomplete question or answer')
    if any(number not in questions for number in short_numbers):
        raise ValueError(f'{sheet.title}: short quiz refers to missing question')
    return {'short_question_numbers': short_numbers,
            'questions': [questions[number] for number in sorted(questions)]}


def knowledge(path):
    book = openpyxl.load_workbook(path, data_only=True)
    items = []
    for sheet in book:
        match = re.fullmatch(r'(Тема|Материал|Тест по теме)\s+([\d.]+)', sheet.title)
        if not match:
            continue
        kind, number = match.groups()
        lines = [clean(sheet.cell(row, 2).value) for row in range(5, sheet.max_row + 1)]
        while lines and not lines[-1]:
            lines.pop()
        answer_key = None
        if kind == 'Тест по теме':
            answer_start = next((index for index, line in enumerate(lines)
                                 if line == 'Ответы и объяснения'), None)
            if answer_start is None:
                raise ValueError(f'{sheet.title}: missing answer section')
            answer_key = '\n\n'.join(str(line) for line in lines[answer_start:] if line)
            lines = lines[:answer_start]
        body = '\n\n'.join(str(line) for line in lines if line)
        title = clean(sheet.cell(2, 2).value)
        sections = sections_from_rows(sheet) if kind != 'Тест по теме' else []
        quiz = quiz_from_sheet(sheet) if kind == 'Тест по теме' else None
        items.append({
            'slug': ('topic-' if kind == 'Тема' else 'material-' if kind == 'Материал' else 'test-') + number.replace('.', '-'),
            'parent_slug': None if kind == 'Тема' else 'topic-' + number.split('.')[0],
            'type': 'topic' if kind == 'Тема' else 'article',
            'kind': 'test' if kind == 'Тест по теме' else 'material' if kind == 'Материал' else 'topic',
            'number': number,
            'title': title,
            'summary': next((line for line in lines if line), '')[:320],
            'body': body,
            'answer_key': answer_key,
            'sections': sections,
            'quiz': quiz,
            'sort_order': int(number.split('.')[0]) * 100 + (int(number.split('.')[1]) if '.' in number else 0) + (90 if kind == 'Тест по теме' else 0),
        })
    assert sum(item['kind'] == 'topic' for item in items) == 4
    assert sum(item['kind'] == 'material' for item in items) == 29
    assert sum(item['kind'] == 'test' for item in items) == 4
    return {'schema_version': 1, 'items': items}


def trainings(path):
    book = openpyxl.load_workbook(path, data_only=True)
    values = {name: table(book[name]) for name in ('Сценарии', 'Флаги и фидбэк', 'Оценка вариантов', 'Тренировки', 'Критерии финала', 'Категории')}
    definitions = []
    for entry in values['Тренировки']:
        training_id = entry['training_id']
        nodes = [row for row in values['Сценарии'] if row['training_id'] == training_id]
        assessments = [row for row in values['Оценка вариантов'] if row['training_id'] == training_id]
        criteria = [row for row in values['Критерии финала'] if row['training_id'] == training_id]
        definitions.append({**entry, 'nodes': nodes, 'assessments': assessments, 'criteria': criteria})
    assert len(definitions) == 8
    assert sum(len(item['nodes']) for item in definitions) == 86
    return {'schema_version': '2.0', 'definitions': definitions,
            'feedback_flags': values['Флаги и фидбэк'], 'categories': values['Категории']}


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for name, data in [('knowledge.json', knowledge(sys.argv[1])),
                       ('trainings.json', trainings(sys.argv[2]))]:
        target = OUTPUT / name
        target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f'{target}: {target.stat().st_size} bytes')


if __name__ == '__main__':
    main()
